"""Native document-generation capability.

Owns Word/PowerPoint/Excel/PDF/text creation, carved out of the
``universal_automation`` god plugin.  The plugin now delegates here, so the
legacy dispatch path is unchanged while this becomes a first-class native
capability the registry can expose directly.

Each generator returns the same result dict shape the plugin returned before,
so no caller behavior changes.  Optional document libraries are imported
lazily; a missing library yields ``{"success": False, "error": ...}``.
"""

from __future__ import annotations

import logging
import os
from typing import Any

from ...utils.file_resolver import resolve_desktop_path
from ..base import ActionSpec, Capability, CapabilityResult, RiskLevel

_WORD_ACTIONS = ("create_word_document", "create_word_doc", "generate_word", "write_word_document")
_PPT_ACTIONS = ("create_powerpoint", "create_ppt", "generate_ppt", "create_presentation")
_EXCEL_ACTIONS = ("create_excel", "create_spreadsheet", "generate_excel", "write_excel")
_PDF_ACTIONS = ("create_pdf", "generate_pdf")
_SAVE_ACTIONS = ("save_to_document", "write_to_file", "append_to_document")


class DocumentCapability(Capability):
    """Create documents (docx, pptx, xlsx, pdf) and plain-text files."""

    name = "documents"
    description = "Create Word, PowerPoint, Excel, PDF, and text documents"
    risk = RiskLevel.MEDIUM

    def discover(self) -> list[ActionSpec]:
        actions = _WORD_ACTIONS + _PPT_ACTIONS + _EXCEL_ACTIONS + _PDF_ACTIONS + _SAVE_ACTIONS
        return [ActionSpec(name=a, risk=RiskLevel.MEDIUM) for a in actions]

    def execute(self, action: str, params: dict[str, Any]) -> CapabilityResult:
        ok, err = self.validate(action, params)
        if not ok:
            return CapabilityResult.fail(err or "validation failed")
        result = self.run(action, params or {})
        return CapabilityResult(
            success=bool(result.get("success")),
            data=result,
            error=result.get("error"),
            metadata={"action": action, "capability": self.name},
            raw=result,
        )

    def run(self, action: str, params: dict[str, Any]) -> dict[str, Any]:
        """Dispatch to the right generator and return the legacy result dict."""
        if action in _WORD_ACTIONS:
            return self.create_word_document(params)
        if action in _PPT_ACTIONS:
            return self.create_powerpoint(params)
        if action in _EXCEL_ACTIONS:
            return self.create_excel(params)
        if action in _PDF_ACTIONS:
            return self.create_pdf(params)
        if action in _SAVE_ACTIONS:
            return self.save_to_document(params)
        return {"success": False, "error": f"Unknown document action '{action}'"}

    # ── generators (behavior copied verbatim from universal_automation) ──────

    def create_word_document(self, params: dict[str, Any]) -> dict[str, Any]:
        """Create a Word document with content"""
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return {
                "success": False,
                "error": "python-docx not installed. Run: pip install python-docx",
            }

        try:
            filename = (
                params.get("filename")
                or params.get("file")
                or params.get("path")
                or "document.docx"
            )
            title = params.get("title", "")
            content = params.get("content") or params.get("text") or params.get("data", "")
            headings = params.get("headings", [])
            folder = params.get("folder") or params.get("directory")

            if folder:
                folder = resolve_desktop_path(folder)
                os.makedirs(folder, exist_ok=True)
                filepath = os.path.join(
                    folder, filename if not filename.startswith("/") else os.path.basename(filename)
                )
            else:
                filepath = resolve_desktop_path(filename)

            if not filepath.endswith(".docx"):
                filepath += ".docx"

            os.makedirs(
                os.path.dirname(filepath) if os.path.dirname(filepath) else ".", exist_ok=True
            )

            doc = Document()

            if title:
                title_para = doc.add_heading(title, 0)
                title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            if isinstance(content, str):
                paragraphs = content.split("\n\n") if "\n\n" in content else content.split("\n")
                for para in paragraphs:
                    if para.strip():
                        if para.strip().isupper() and len(para.strip()) < 100:
                            doc.add_heading(para.strip(), level=2)
                        else:
                            doc.add_paragraph(para.strip())

            elif isinstance(content, dict):
                for key, value in content.items():
                    doc.add_heading(str(key), level=2)
                    if isinstance(value, list):
                        for item in value:
                            doc.add_paragraph(str(item), style="List Bullet")
                    else:
                        doc.add_paragraph(str(value))

            elif isinstance(content, list):
                for item in content:
                    if isinstance(item, dict):
                        for k, v in item.items():
                            doc.add_heading(str(k), level=2)
                            doc.add_paragraph(str(v))
                    else:
                        doc.add_paragraph(str(item))

            if headings:
                for heading in headings:
                    if isinstance(heading, dict):
                        h_title = heading.get("title", "")
                        h_content = heading.get("content", "")
                        h_level = heading.get("level", 2)
                        doc.add_heading(h_title, level=h_level)
                        if h_content:
                            doc.add_paragraph(h_content)

            doc.save(filepath)

            return {
                "success": True,
                "message": "Word document created successfully",
                "filepath": filepath,
                "filename": os.path.basename(filepath),
            }

        except Exception as e:
            logging.getLogger(__name__).exception("Word document creation failed")
            return {"success": False, "error": str(e)}

    def create_powerpoint(self, params: dict[str, Any]) -> dict[str, Any]:
        """Create a PowerPoint presentation"""
        try:
            from pptx import Presentation
        except ImportError:
            return {
                "success": False,
                "error": "python-pptx not installed. Run: pip install python-pptx",
            }

        try:
            filename = (
                params.get("filename")
                or params.get("file")
                or params.get("path")
                or "presentation.pptx"
            )
            title = params.get("title", "Presentation")
            slides_data = params.get("slides") or params.get("content", [])
            folder = params.get("folder") or params.get("directory")

            if folder:
                folder = resolve_desktop_path(folder)
                os.makedirs(folder, exist_ok=True)
                filepath = os.path.join(
                    folder, filename if not filename.startswith("/") else os.path.basename(filename)
                )
            else:
                filepath = resolve_desktop_path(filename)

            if not filepath.endswith(".pptx"):
                filepath += ".pptx"

            os.makedirs(
                os.path.dirname(filepath) if os.path.dirname(filepath) else ".", exist_ok=True
            )

            prs = Presentation()

            title_slide_layout = prs.slide_layouts[0]
            slide = prs.slides.add_slide(title_slide_layout)
            slide.shapes.title.text = title
            if params.get("subtitle"):
                slide.placeholders[1].text = params.get("subtitle")

            if isinstance(slides_data, str):
                sections = slides_data.split("\n\n")
                for section in sections:
                    if section.strip():
                        bullet_slide_layout = prs.slide_layouts[1]
                        slide = prs.slides.add_slide(bullet_slide_layout)
                        lines = section.split("\n")
                        slide.shapes.title.text = lines[0][:100] if lines else "Content"
                        if len(lines) > 1:
                            text_frame = slide.placeholders[1].text_frame
                            for line in lines[1:]:
                                if line.strip():
                                    p = text_frame.add_paragraph()
                                    p.text = line.strip()
                                    p.level = 0

            elif isinstance(slides_data, list):
                for slide_content in slides_data:
                    bullet_slide_layout = prs.slide_layouts[1]
                    slide = prs.slides.add_slide(bullet_slide_layout)

                    if isinstance(slide_content, dict):
                        slide.shapes.title.text = slide_content.get("title", "Slide")
                        content = slide_content.get("content", [])

                        if isinstance(content, list):
                            text_frame = slide.placeholders[1].text_frame
                            text_frame.clear()
                            for item in content:
                                p = text_frame.add_paragraph()
                                p.text = str(item)
                                p.level = 0
                        elif isinstance(content, str):
                            slide.placeholders[1].text = content
                    else:
                        slide.shapes.title.text = str(slide_content)[:100]

            prs.save(filepath)

            return {
                "success": True,
                "message": "PowerPoint presentation created successfully",
                "filepath": filepath,
                "filename": os.path.basename(filepath),
                "slides_count": len(prs.slides),
            }

        except Exception as e:
            logging.getLogger(__name__).exception("PowerPoint creation failed")
            return {"success": False, "error": str(e)}

    def create_excel(self, params: dict[str, Any]) -> dict[str, Any]:
        """Create an Excel spreadsheet"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError:
            return {"success": False, "error": "openpyxl not installed. Run: pip install openpyxl"}

        try:
            filename = (
                params.get("filename")
                or params.get("file")
                or params.get("path")
                or "spreadsheet.xlsx"
            )
            data = params.get("data") or params.get("content", [])
            headers = params.get("headers", [])
            sheet_name = params.get("sheet_name", "Sheet1")
            folder = params.get("folder") or params.get("directory")

            if folder:
                folder = resolve_desktop_path(folder)
                os.makedirs(folder, exist_ok=True)
                filepath = os.path.join(
                    folder, filename if not filename.startswith("/") else os.path.basename(filename)
                )
            else:
                filepath = resolve_desktop_path(filename)

            if not filepath.endswith(".xlsx"):
                filepath += ".xlsx"

            os.makedirs(
                os.path.dirname(filepath) if os.path.dirname(filepath) else ".", exist_ok=True
            )

            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name

            if headers:
                for col, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col, value=header)
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill(
                        start_color="4472C4", end_color="4472C4", fill_type="solid"
                    )
                    cell.alignment = Alignment(horizontal="center")

            start_row = 2 if headers else 1

            if isinstance(data, list):
                for row_idx, row_data in enumerate(data, start=start_row):
                    if isinstance(row_data, list):
                        for col_idx, value in enumerate(row_data, start=1):
                            ws.cell(row=row_idx, column=col_idx, value=value)
                    elif isinstance(row_data, dict):
                        for col_idx, key in enumerate(
                            headers if headers else row_data.keys(), start=1
                        ):
                            ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))
                    else:
                        ws.cell(row=row_idx, column=1, value=str(row_data))

            elif isinstance(data, dict):
                for row_idx, (key, value) in enumerate(data.items(), start=start_row):
                    ws.cell(row=row_idx, column=1, value=key)
                    if isinstance(value, list):
                        for col_idx, item in enumerate(value, start=2):
                            ws.cell(row=row_idx, column=col_idx, value=item)
                    else:
                        ws.cell(row=row_idx, column=2, value=value)

            for column in ws.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            wb.save(filepath)

            return {
                "success": True,
                "message": "Excel spreadsheet created successfully",
                "filepath": filepath,
                "filename": os.path.basename(filepath),
                "rows": len(data)
                if isinstance(data, list)
                else len(data.keys())
                if isinstance(data, dict)
                else 0,
            }

        except Exception as e:
            logging.getLogger(__name__).exception("Excel creation failed")
            return {"success": False, "error": str(e)}

    def create_pdf(self, params: dict[str, Any]) -> dict[str, Any]:
        """Create a PDF document"""
        try:
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import inch
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError:
            return {
                "success": False,
                "error": "reportlab not installed. Run: pip install reportlab",
            }

        try:
            filename = (
                params.get("filename") or params.get("file") or params.get("path") or "document.pdf"
            )
            title = params.get("title", "")
            content = params.get("content") or params.get("text") or params.get("data", "")
            folder = params.get("folder") or params.get("directory")

            if folder:
                folder = resolve_desktop_path(folder)
                os.makedirs(folder, exist_ok=True)
                filepath = os.path.join(
                    folder, filename if not filename.startswith("/") else os.path.basename(filename)
                )
            else:
                filepath = resolve_desktop_path(filename)

            if not filepath.endswith(".pdf"):
                filepath += ".pdf"

            os.makedirs(
                os.path.dirname(filepath) if os.path.dirname(filepath) else ".", exist_ok=True
            )

            doc = SimpleDocTemplate(filepath, pagesize=letter)
            story = []
            styles = getSampleStyleSheet()

            if title:
                title_style = ParagraphStyle(
                    "CustomTitle",
                    parent=styles["Heading1"],
                    fontSize=24,
                    textColor="#1f4788",
                    spaceAfter=30,
                    alignment=TA_CENTER,
                )
                story.append(Paragraph(title, title_style))
                story.append(Spacer(1, 0.2 * inch))

            if isinstance(content, str):
                paragraphs = content.split("\n\n") if "\n\n" in content else content.split("\n")
                for para in paragraphs:
                    if para.strip():
                        story.append(Paragraph(para.strip(), styles["BodyText"]))
                        story.append(Spacer(1, 0.1 * inch))

            elif isinstance(content, list):
                for item in content:
                    story.append(Paragraph(str(item), styles["BodyText"]))
                    story.append(Spacer(1, 0.1 * inch))

            doc.build(story)

            return {
                "success": True,
                "message": "PDF document created successfully",
                "filepath": filepath,
                "filename": os.path.basename(filepath),
            }

        except Exception as e:
            logging.getLogger(__name__).exception("PDF creation failed")
            return {"success": False, "error": str(e)}

    def save_to_document(self, params: dict[str, Any]) -> dict[str, Any]:
        """Save content to a document (auto-detect format or append)"""
        try:
            filepath = params.get("file") or params.get("path") or params.get("filename")
            content = params.get("content") or params.get("text") or params.get("data", "")
            append = params.get("append", False)

            if not filepath:
                return {"success": False, "error": "No filepath provided"}

            filepath = resolve_desktop_path(filepath)

            if filepath.endswith(".docx"):
                if append and os.path.exists(filepath):
                    try:
                        from docx import Document

                        doc = Document(filepath)
                        if isinstance(content, str):
                            for para in content.split("\n"):
                                if para.strip():
                                    doc.add_paragraph(para.strip())
                        doc.save(filepath)
                        return {"success": True, "filepath": filepath, "mode": "append"}
                    except Exception:
                        pass
                return self.create_word_document({**params, "filename": filepath, "content": content})

            elif filepath.endswith(".pptx"):
                return self.create_powerpoint({**params, "filename": filepath})

            elif filepath.endswith(".xlsx"):
                return self.create_excel({**params, "filename": filepath})

            elif filepath.endswith(".pdf"):
                return self.create_pdf({**params, "filename": filepath})

            else:
                os.makedirs(
                    os.path.dirname(filepath) if os.path.dirname(filepath) else ".", exist_ok=True
                )
                mode = "a" if append else "w"
                with open(filepath, mode, encoding="utf-8") as f:
                    f.write(str(content))
                    if not str(content).endswith("\n"):
                        f.write("\n")

                return {
                    "success": True,
                    "filepath": filepath,
                    "mode": "append" if append else "write",
                }

        except Exception as e:
            return {"success": False, "error": str(e)}
